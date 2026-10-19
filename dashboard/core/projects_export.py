"""SI extract for the Projects scope (Seq 22).

One workbook the System Integrator can work from to load the new ERP's
projects module: project headers, tasks (HAIS codes), the opening position
per project and per task, the year-by-year history, and a reconciliation.

Used by the Projects tab's download button and by
scripts/generate_projects_export.py, so both produce the same file.
"""
import io
from datetime import date

import pandas as pd

from dashboard.core.projects import (
    COLUMN_LABELS, HOUSES, LIVE_MONTHS, activity_signal, coverage, fy_label, opening_table, period_label,
    scope_view, unit_history,
)

_HEADER_FILL = '1E1528'


def _filters_text(include_annual, expenditure_only):
    return ('All listed codes, including Annual Requirements' if include_annual else
            'Projects only (CAT A, CAT B, Other), Annual Requirements excluded') + \
           ('; expenditure only (income and recharge cost types excluded, provisional rule)'
            if expenditure_only else '; all cost types')


def build_sheets(frames: dict, include_annual=True, expenditure_only=False) -> dict:
    view = scope_view(frames, include_annual, expenditure_only)
    table = opening_table(view)
    codes = view['codes'].drop_duplicates('hais_code')
    bounds = view['bounds']
    lines = view['lines']

    header = table[['unit', 'unit_desc', 'project_type', 'projgrp_desc', 'directorate', 'houses', 'signal',
                    'first_activity', 'last_activity', 'hais_listed']].merge(
        view['listed_units'][['unit', 'projcat', 'projgrp', 'costc', 'costc_desc', 'directorate_desc',
                              'department', 'department_desc']], on='unit', how='left')
    header['active_lines'] = header['unit'].map(lines.groupby('unit')['hais_code'].nunique()).fillna(0).astype(int)
    for gap in ('Project manager', 'Project start date', 'Project end date'):
        header[gap] = ''
    header = header.rename(columns={**COLUMN_LABELS, 'active_lines': 'Cost lines with activity'})

    tasks = codes[['hais_code', 'hais_desc', 'unit', 'unit_desc', 'project_type', 'status', 'houses',
                   'first_period', 'last_period', 'has_activity']].copy()
    tasks['signal'] = activity_signal(tasks['last_period'], bounds)
    tasks['first_period'] = tasks['first_period'].map(period_label)
    tasks['last_period'] = tasks['last_period'].map(period_label)
    tasks = tasks.rename(columns={**COLUMN_LABELS, 'first_period': 'First activity', 'last_period': 'Last activity'})

    pos = view['positions'][['unit', 'unit_desc', 'project_type', 'house', 'spend_to_date', 'commitments',
                             'budget_to_date', 'remaining_budget', 'remaining_after_commitments',
                             'current_fy_actuals', 'current_fy_budget']].round(2)
    pos = pos.sort_values(['unit', 'house']).rename(columns=COLUMN_LABELS)

    task_pos = lines.groupby(['unit', 'hais_code', 'hais_desc', 'house']).agg(
        spend_to_date=('actuals', 'sum'), commitments=('committed', 'sum'), budget_to_date=('budget', 'sum')
    ).reset_index()
    task_pos['remaining_budget'] = task_pos['budget_to_date'] - task_pos['spend_to_date']
    task_pos = task_pos.round(2).rename(columns=COLUMN_LABELS)

    hist = unit_history(lines)[['unit', 'house', 'fy_label', 'actuals', 'cumulative_actuals', 'budget',
                                'cumulative_budget', 'remaining_budget', 'committed']].round(2)
    hist = hist.rename(columns={**COLUMN_LABELS, 'committed': 'Open commitments posted to this year'})

    cov = coverage(frames)
    recon = pd.DataFrame([
        ('Cost lines on the project list', cov['listed_codes']),
        ('Cost lines in the financials', cov['fin_codes']),
        ('Listed cost lines with activity', cov['matched']),
        ('Listed cost lines with no activity since April 2020', cov['idle']),
        ('Cost lines in the financials but not on the list (excluded)', cov['off_list']),
        ('Actuals on listed cost lines (raw file)', round(cov['on_list_actuals'], 2)),
        ('Spend to date across all project positions', round(cov['position_actuals'], 2)),
        ('Positions reconcile to the raw file', 'Yes' if cov['invariant_ok'] else 'NO - investigate'),
        ('This extract (after filters): spend to date', round(float(view['positions']['spend_to_date'].sum()), 2)),
        ('This extract (after filters): open commitments', round(float(view['positions']['commitments'].sum()), 2)),
        ('This extract (after filters): budget to date', round(float(view['positions']['budget_to_date'].sum()), 2)),
    ], columns=['Check', 'Value'])

    readme = pd.DataFrame([
        ('Scope', 'Seq 22: opening cumulative position (spend to date, commitments and remaining budget) '
                  'for all projects active at go-live.'),
        ('Produced', date.today().isoformat()),
        ('Data as at', period_label(bounds.get('latest'))),
        ('History from', period_label(bounds.get('first'))),
        ('Filters applied', _filters_text(include_annual, expenditure_only)),
        ('Project', 'A project is a Unit. Its HAIS codes are the cost lines (tasks) under it.'),
        ('Spend to date', 'Sum of actuals since the start of the data.'),
        ('Open commitments', 'Sum of committed purchase: the amount left on purchase orders.'),
        ('Budget to date', 'Sum of current budget across every year, rolled up to the project.'),
        ('Remaining budget', 'Budget to date minus spend to date.'),
        ('Activity signal', f'Live = activity in the last {LIVE_MONTHS} months; Dormant = older activity only; '
                            'No activity = nothing since April 2020.'),
        ('Houses', 'Positions are held per house. A shared cost line has its amounts already split by house '
                   'in the source (Recharge).'),
        ('Known gap', 'Project manager and start / end dates are not held in either source file. Columns are '
                      'provided blank on Project_Header.'),
        ('Known gap', 'Financials start in April 2020. Projects active then may have earlier spend that is not '
                      'included (see DQ check PRJ_HISTORY_TRUNCATED).'),
        ('Sources', 'Project list (Finance, by HAIS code) and Finance raw data dump (actuals, budget, '
                    'committed purchase by account, HAIS code, house and period).'),
    ], columns=['Item', 'Detail'])

    return {'README': readme, 'Project_Header': header, 'Project_Tasks': tasks, 'Opening_Position': pos,
            'Task_Position': task_pos, 'Annual_History': hist, 'Reconciliation': recon}


def write_workbook(sheets: dict, target) -> None:
    """Write sheets to a path or a file-like object, with readable formatting."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    with pd.ExcelWriter(target, engine='openpyxl') as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            for cell in ws[1]:
                cell.fill = PatternFill('solid', fgColor=_HEADER_FILL)
                cell.font = Font(color='FFFFFF', bold=True)
                cell.alignment = Alignment(vertical='center', wrap_text=True)
            ws.freeze_panes = 'A2'
            for i, col in enumerate(df.columns, start=1):
                width = max([len(str(col))] + [len(str(v)) for v in df[col].head(200).tolist()])
                ws.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 70)
                if pd.api.types.is_float_dtype(df[col]):
                    for row in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                        row[0].number_format = '#,##0.00'
            if name == 'README':
                ws.column_dimensions['B'].width = 110
                for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
                    row[0].alignment = Alignment(wrap_text=True, vertical='top')


def workbook_bytes(frames: dict, include_annual=True, expenditure_only=False) -> bytes:
    buf = io.BytesIO()
    write_workbook(build_sheets(frames, include_annual, expenditure_only), buf)
    return buf.getvalue()
